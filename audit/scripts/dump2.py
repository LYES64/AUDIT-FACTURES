import openpyxl
def dump(f, sheet):
    wb = openpyxl.load_workbook(f, data_only=True)
    ws = wb[sheet]
    print("###### ",sheet, " max_col=",ws.max_column)
    for i,r in enumerate(ws.iter_rows(values_only=True)):
        if any(c is not None and str(c).strip()!="" for c in r):
            print(i,"|".join("" if c is None else str(c) for c in r))
f="/root/.claude/uploads/7767e6b5-de9b-5631-a5c3-b77beef4874c/b7dfb6a3-Suivi_IRSH_2026.xlsx"
dump(f,'Copie de non facturé')
