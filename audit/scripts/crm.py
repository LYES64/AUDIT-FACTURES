import openpyxl
f="/root/.claude/uploads/7767e6b5-de9b-5631-a5c3-b77beef4874c/ae2485a5-Compiexe__Ind_pendance_Royale_S_nior_Habitat_1.xlsx"
wb=openpyxl.load_workbook(f,data_only=True)
print("Feuilles:",wb.sheetnames)
for ws in wb.worksheets:
    print(f"\n=== {ws.title}  ({ws.max_row}x{ws.max_column}) ===")
    for i,r in enumerate(ws.iter_rows(values_only=True)):
        if i>6: break
        print(i,[ ("" if c is None else str(c))[:22] for c in r])
