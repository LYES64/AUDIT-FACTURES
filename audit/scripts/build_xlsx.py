import csv,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
wb=openpyxl.Workbook()
hdr=Font(bold=True,color="FFFFFF"); fill=PatternFill("solid",fgColor="1F4E78")
def stylehdr(ws):
    for c in ws[1]:
        c.font=hdr; c.fill=fill; c.alignment=Alignment(wrap_text=True,vertical="center")
    ws.freeze_panes="A2"

# Sheet 1: Synthèse
ws=wb.active; ws.title="Synthèse Dossiers 2026"
rows=list(csv.reader(open("Audit_Factures_2026_Synthese.csv")))
for r in rows: ws.append(r)
stylehdr(ws)
widths=[16,16,32,14,12,22,14,12,12,26,26,12,42,14]
for i,w in enumerate(widths,1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
# color rows by anomaly
warn=PatternFill("solid",fgColor="FFF2CC"); red=PatternFill("solid",fgColor="F8CBAD")
for ri in range(2,ws.max_row+1):
    st=ws.cell(ri,13).value or ""
    if "Anomalie" in st: 
        for ci in range(1,15): ws.cell(ri,ci).fill=red
    elif "À facturer" in st or "À vérifier" in st:
        for ci in range(1,15): ws.cell(ri,ci).fill=warn

# Sheet 2: Anomalies
ws2=wb.create_sheet("Anomalies")
anoms=[
["#","Prescripteur","Anomalie","Dossier / Réf","Preuve (fichiers)","Pourquoi c'est une anomalie","Urgence"],
[1,"IRSH","9 factures avec MONTANT mais N° EFFACÉ dans 'Suivi Factures'","Lignes 47-55 (2460,2100,5010,2450,1800,3300,2950,4890,2700€) = onglets F2026-16 à F2026-24","Onglets F2026-16..24 datés 10-12/06/2026 ; montants identiques aux lignes sans numéro","Les numéros de facture ont été retirés du suivi alors que les factures existent (onglets). Suppression/masquage manifeste.","CRITIQUE"],
[2,"IRSH","Séquence d'avoirs incomplète : AV03, AV04, AV05 absents","AV01,AV02 puis saut à AV06..AV14","Suivi Factures IRSH","Trou dans la numérotation des avoirs — 3 avoirs potentiellement supprimés ou jamais tracés.","ÉLEVÉE"],
[3,"IRSH","Écart de montant onglet vs suivi : F2026-23","Onglet=4850€ / ligne suivi=4890€ (Δ+40€)","Onglet F2026-23 vs Suivi ligne 54","Montant modifié entre la facture et le suivi.","MOYENNE"],
[4,"IRSH","F2026-15 orpheline : 200€, sans date, sans état, sans réglement","F2026-15","Suivi Factures ligne 45","Facture incomplète / statut indéterminé.","MOYENNE"],
[5,"IRSH","21 poses réalisées ('OUI', validées) SANS N° de facture dans le suivi","cf. onglet Suivi Poses","Suivi Poses IRSH lignes 9-35","Poses faites mais non rattachées à une facture dans les fichiers — factures potentiellement jamais émises (à confirmer par mails).","ÉLEVÉE"],
[6,"IRSH","Liste 'Copie de non facturé' = 30 170€ HT de poses non facturées","onglet 'Copie de non facturé'","36 lignes DIA avec montants","Volume important de CA potentiellement non facturé.","ÉLEVÉE"],
[7,"BARROS","Onglets F2026-B11 et F2026-B12 ABSENTS du 'Suivi Factures'","B11 (4750€), B12 (0€, brouillon)","Onglets présents, suivi s'arrête à B10","Factures générées mais non enregistrées dans le suivi.","ÉLEVÉE"],
[8,"BARROS","Ligne sans numéro : 7600€","Suivi Factures ligne 24","Montant 7600€ sans N°","Facture avec montant mais numéro absent — ne correspond ni à B11 ni à B12.","ÉLEVÉE"],
[9,"BARROS","N° internes erronés : onglets B11 ET B12 affichent 'F2026-B10'","B11, B12","Cellule 'Facture N°' des onglets","Copier-coller sans mise à jour du numéro — risque de doublon de numérotation.","MOYENNE"],
[10,"BARROS","Écart montant onglet vs suivi : F2026-B01","Onglet=13580€ / suivi=13850€ (Δ+270€)","Onglet B01 vs Suivi","Montant divergent entre facture et suivi.","MOYENNE"],
[11,"BARROS","33 commandes CO-xxx marquées facturées mais tracées comme SAV/sans détail","Suivi Poses BARROS","Poses B02-B08 sans onglet de détail","Impossible de vérifier le contenu facturé sans les factures/mails.","MOYENNE"],
[12,"2 fichiers","Numérotation 2025 datée 13/01/2026 (BARROS) et 15/01/2026 (IRSH)","F0925,F1025,F1125,F1225,F122025 / F1225-xx","Suivi Factures","Factures de fin 2025 saisies/datées en janvier 2026 — antidatage possible, à clarifier.","MOYENNE"],
]
for r in anoms: ws2.append(r)
stylehdr(ws2)
for i,w in enumerate([4,12,46,40,40,50,12],1): ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
for ri in range(2,ws2.max_row+1):
    ws2.cell(ri,3).alignment=Alignment(wrap_text=True,vertical="top")
    for ci in (5,6): ws2.cell(ri,ci).alignment=Alignment(wrap_text=True,vertical="top")
    u=ws2.cell(ri,7).value
    ws2.cell(ri,7).fill=PatternFill("solid",fgColor="F8CBAD" if u=="CRITIQUE" else "FCE4D6" if u=="ÉLEVÉE" else "FFF2CC")

wb.save("Audit_Factures_2026.xlsx")
print("xlsx OK")
