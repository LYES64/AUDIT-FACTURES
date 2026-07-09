import openpyxl
for f,label in [("/root/.claude/uploads/7767e6b5-de9b-5631-a5c3-b77beef4874c/b7dfb6a3-Suivi_IRSH_2026.xlsx","IRSH"),
                ("/root/.claude/uploads/7767e6b5-de9b-5631-a5c3-b77beef4874c/241efaeb-Suivi_Barros_2026.xlsx","BARROS")]:
    wb = openpyxl.load_workbook(f, data_only=True)
    print("="*70)
    print("FICHIER:", label, "  Feuilles:", wb.sheetnames)
    for ws in wb.worksheets:
        print(f"  -- '{ws.title}': max_row={ws.max_row} max_col={ws.max_column}")
