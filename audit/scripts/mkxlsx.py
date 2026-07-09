import csv,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
rows=list(csv.reader(open("IRSH_Interventions_a_facturer.csv")))
wb=openpyxl.Workbook(); ws=wb.active; ws.title="IRSH - A facturer (CRM 07-07)"
for r in rows: ws.append(r)
for c in ws[1]:
    c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F4E78")
    c.alignment=Alignment(wrap_text=True,vertical="center")
ws.freeze_panes="A2"
for i,w in enumerate([46,12,8,32,14,16,16,12],1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
colors={"1":"C6EFCE","2":"FFF2CC","3":"DDEBF7","4":"F2F2F2"}
for ri in range(2,ws.max_row+1):
    cat=ws.cell(ri,1).value or ""
    key=cat.strip()[0] if cat else ""
    fill=colors.get(key)
    if fill:
        for ci in range(1,9): ws.cell(ri,ci).fill=PatternFill("solid",fgColor=fill)
wb.save("IRSH_Interventions_a_facturer.xlsx")
print("ok")
